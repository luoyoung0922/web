# -*- coding: utf-8 -*-
import json
import os
import ssl
import sqlite3
import uuid
from datetime import datetime
from functools import wraps
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import pdfplumber
import certifi
from docx import Document
from flask import Flask, abort, jsonify, render_template, request, send_from_directory, session
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

try:
    from openai import OpenAI
except Exception:
    OpenAI = None


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "data.sqlite3")
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
FIXED_ADMIN_USERNAME = os.environ.get("FIXED_ADMIN_USER", "admin")
FIXED_ADMIN_PASSWORD = os.environ.get("FIXED_ADMIN_PASSWORD", "admin123456")
DEFAULT_AI_API_KEY = os.environ.get("AI_API_KEY") or os.environ.get("OPENAI_API_KEY") or ""
DEFAULT_AI_API_BASE = os.environ.get("AI_API_BASE") or os.environ.get("OPENAI_BASE_URL") or ""
DEFAULT_AI_MODEL = os.environ.get("AI_MODEL") or os.environ.get("OPENAI_MODEL") or "gpt-4o-mini"
DEFAULT_AI_API_TYPE = os.environ.get("AI_API_TYPE", "chat")

app = Flask(__name__, static_folder="static", template_folder="templates")
app.secret_key = os.environ.get("SECRET_KEY", "dev-secret-change-me")
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024
os.makedirs(UPLOAD_DIR, exist_ok=True)


def now_iso():
    return datetime.utcnow().isoformat()


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def ensure_column(db, table, column, definition):
    cols = [row["name"] for row in db.execute(f"PRAGMA table_info({table})").fetchall()]
    if column not in cols:
        db.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def init_db():
    with get_db() as db:
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS admins (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                is_owner INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS papers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                file_type TEXT NOT NULL,
                filename TEXT NOT NULL,
                stored_filename TEXT NOT NULL DEFAULT '',
                raw_text TEXT NOT NULL,
                use_ai INTEGER NOT NULL DEFAULT 1,
                created_by INTEGER,
                created_at TEXT NOT NULL
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS questions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                paper_id INTEGER NOT NULL,
                qtype TEXT NOT NULL DEFAULT 'fill',
                prompt TEXT NOT NULL,
                answer TEXT NOT NULL,
                choices_json TEXT NOT NULL DEFAULT '[]',
                explanation TEXT NOT NULL DEFAULT '',
                full_text TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(paper_id) REFERENCES papers(id) ON DELETE CASCADE
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS attempts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                question_id INTEGER NOT NULL,
                user_id INTEGER NOT NULL,
                user_answer TEXT NOT NULL,
                is_correct INTEGER NOT NULL,
                score REAL NOT NULL DEFAULT 0,
                feedback TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                FOREIGN KEY(question_id) REFERENCES questions(id) ON DELETE CASCADE
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS paper_participants (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                paper_id INTEGER NOT NULL,
                user_id INTEGER NOT NULL,
                participate INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(paper_id, user_id),
                FOREIGN KEY(paper_id) REFERENCES papers(id) ON DELETE CASCADE
            )
            """
        )
        ensure_column(db, "admins", "is_owner", "INTEGER NOT NULL DEFAULT 0")
        ensure_column(db, "papers", "file_type", "TEXT NOT NULL DEFAULT '资料'")
        ensure_column(db, "papers", "use_ai", "INTEGER NOT NULL DEFAULT 1")
        ensure_column(db, "papers", "created_by", "INTEGER")
        ensure_column(db, "questions", "paper_id", "INTEGER NOT NULL DEFAULT 0")
        ensure_column(db, "attempts", "user_id", "INTEGER NOT NULL DEFAULT 0")
        ensure_column(db, "attempts", "participate", "INTEGER NOT NULL DEFAULT 0")
        db.commit()


def get_setting(name, default=""):
    with get_db() as db:
        row = db.execute("SELECT value FROM settings WHERE key = ?", (name,)).fetchone()
    return row["value"] if row else default


def set_setting(name, value):
    with get_db() as db:
        db.execute(
            "INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value",
            (name, str(value)),
        )
        db.commit()


def bootstrap_fixed_admin():
    with get_db() as db:
        owner_id = get_setting("owner_admin_id", "").strip()
        if owner_id.isdigit():
            row = db.execute("SELECT id FROM admins WHERE id = ?", (int(owner_id),)).fetchone()
            if row:
                db.execute("UPDATE admins SET is_owner = CASE WHEN id = ? THEN 1 ELSE 0 END", (int(owner_id),))
                db.commit()
                return

        row = db.execute("SELECT id FROM admins WHERE username = ?", (FIXED_ADMIN_USERNAME,)).fetchone()
        if row:
            admin_id = row["id"]
            db.execute("UPDATE admins SET is_owner = CASE WHEN id = ? THEN 1 ELSE 0 END", (admin_id,))
        else:
            admin_id = db.execute(
                "INSERT INTO admins (username, password_hash, is_owner, created_at) VALUES (?, ?, 1, ?)",
                (FIXED_ADMIN_USERNAME, generate_password_hash(FIXED_ADMIN_PASSWORD), now_iso()),
            ).lastrowid
            db.execute("UPDATE admins SET is_owner = CASE WHEN id = ? THEN 1 ELSE 0 END", (admin_id,))

        for key, value in [
            ("owner_admin_id", str(admin_id)),
            ("ai_api_key", DEFAULT_AI_API_KEY),
            ("ai_api_base", DEFAULT_AI_API_BASE),
            ("ai_model", DEFAULT_AI_MODEL),
            ("ai_api_type", DEFAULT_AI_API_TYPE),
        ]:
            db.execute(
                "INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO NOTHING",
                (key, value),
            )
        db.commit()


def current_identity():
    if session.get("admin_id"):
        return {
            "kind": "admin",
            "id": session.get("admin_id"),
            "username": session.get("admin_username"),
            "is_owner": bool(session.get("is_owner")),
        }
    if session.get("user_id"):
        return {
            "kind": "user",
            "id": session.get("user_id"),
            "username": session.get("user_username"),
        }
    return None


def login_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if not session.get("user_id") and not session.get("admin_id"):
            return jsonify({"error": "login_required"}), 401
        return view(*args, **kwargs)

    return wrapper


def admin_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if not session.get("admin_id"):
            return jsonify({"error": "unauthorized"}), 401
        return view(*args, **kwargs)

    return wrapper


def owner_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if not session.get("admin_id"):
            return jsonify({"error": "unauthorized"}), 401
        if not session.get("is_owner"):
            return jsonify({"error": "owner_required"}), 403
        return view(*args, **kwargs)

    return wrapper


def normalize(text):
    text = str(text).lower()
    for ch in " ，。！？、,.!?;；:：\n\t\r":
        text = text.replace(ch, "")
    return text


def file_type_label(filename):
    ext = os.path.splitext(filename)[1].lower()
    return {
        ".pdf": "PDF",
        ".docx": "Word",
        ".txt": "文本",
        ".text": "文本",
        ".md": "Markdown",
        ".markdown": "Markdown",
    }.get(ext, "资料")


def parse_pdf_text(path):
    texts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            texts.append(page.extract_text() or "")
    return "\n".join(texts)


def parse_docx_text(path):
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def parse_text_file(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def parse_uploaded_file(path, original_name):
    ext = os.path.splitext(original_name)[1].lower()
    if ext == ".pdf":
        return parse_pdf_text(path)
    if ext in {".txt", ".text", ".md", ".markdown"}:
        return parse_text_file(path)
    if ext == ".docx":
        return parse_docx_text(path)
    raise ValueError("unsupported_file_type")


def save_uploaded_file(file):
    original = secure_filename(file.filename) or "document"
    stored_filename = f"{uuid.uuid4().hex}_{original}"
    path = os.path.join(UPLOAD_DIR, stored_filename)
    file.save(path)
    return stored_filename, path


def split_sentences(text):
    normalized = text.replace("\r", " ").replace("\n", "。")
    parts = []
    for part in normalized.split("。"):
        s = part.strip()
        if len(s) > 18:
            parts.append(s)
    return parts[:200]


def fallback_build_question_bank(raw_text):
    items = []
    for sentence in split_sentences(raw_text)[:50]:
        words = [w for w in sentence.replace("，", " ").replace("、", " ").split() if w]
        answer = max(words, key=len) if words else sentence[:8]
        prompt = sentence.replace(answer, "______", 1)
        items.append(
            {
                "type": "fill",
                "prompt": prompt,
                "answer": answer,
                "choices": [],
                "explanation": "来自原文的自动填空题。",
                "source_text": sentence,
            }
        )
    return items


def normalize_question_item(item, source_text=""):
    qtype = str(item.get("type", "fill")).strip().lower()
    prompt = str(item.get("prompt", "")).strip()
    answer = str(item.get("answer", "")).strip()
    choices = item.get("choices") or []
    explanation = str(item.get("explanation", "")).strip()
    source_text = str(item.get("source_text", "") or source_text).strip()

    if qtype not in {"choice", "judge", "fill", "short"}:
        qtype = "fill"

    if qtype == "judge":
        n = normalize(answer)
        if n in {"对", "正确", "true", "yes", "t", "1"}:
            answer = "对"
        elif n in {"错", "错误", "false", "no", "f", "0"}:
            answer = "错"
        else:
            answer = "对" if "对" in n or "正确" in n else "错"
        choices = ["对", "错"]
    elif qtype == "choice":
        cleaned = []
        for choice in choices:
            text = str(choice).strip()
            if text and text not in cleaned:
                cleaned.append(text)
        if len(cleaned) < 2:
            cleaned = []
        if len(cleaned) > 4:
            cleaned = cleaned[:4]
        if cleaned and answer and answer not in cleaned:
            cleaned[0] = answer
        choices = cleaned
    else:
        choices = []

    if not prompt:
        prompt = source_text or answer or "题目"
    if not answer:
        answer = source_text[:20] if source_text else "答案"

    return {
        "type": qtype,
        "prompt": prompt,
        "answer": answer,
        "choices": choices,
        "explanation": explanation,
        "source_text": source_text or prompt,
    }


def ai_client():
    api_key = get_setting("ai_api_key", DEFAULT_AI_API_KEY)
    if not api_key or OpenAI is None:
        return None
    kwargs = {"api_key": api_key}
    base_url = get_setting("ai_api_base", DEFAULT_AI_API_BASE)
    if base_url:
        kwargs["base_url"] = base_url
    return OpenAI(**kwargs)


def ai_api_key():
    return get_setting("ai_api_key", DEFAULT_AI_API_KEY)


def ai_api_base_url():
    return (get_setting("ai_api_base", DEFAULT_AI_API_BASE) or "https://api.openai.com/v1").strip().rstrip("/")


def ai_model():
    return get_setting("ai_model", DEFAULT_AI_MODEL)


def ai_api_type():
    value = (get_setting("ai_api_type", DEFAULT_AI_API_TYPE) or "chat").strip().lower()
    return "responses" if value == "responses" else "chat"


def ai_enabled():
    return bool(ai_api_key())


def parse_json_content(content):
    text = (content or "{}").strip()
    if text.startswith("```"):
        text = text.strip("`").strip()
        if text.lower().startswith("json"):
            text = text[4:].strip()
    data = json.loads(text or "{}")
    return data if isinstance(data, dict) else None


def response_output_text(response):
    if isinstance(response, dict):
        if response.get("output_text"):
            return response["output_text"]
        chunks = []
        for item in response.get("output", []) or []:
            for content in item.get("content", []) or []:
                text = content.get("text") if isinstance(content, dict) else None
                if text:
                    chunks.append(text)
        return "".join(chunks)

    direct_text = getattr(response, "output_text", None)
    if direct_text:
        return direct_text
    chunks = []
    for item in getattr(response, "output", []) or []:
        for content in getattr(item, "content", []) or []:
            text = getattr(content, "text", None)
            if text:
                chunks.append(text)
    return "".join(chunks)


def responses_endpoint_url():
    base_url = ai_api_base_url()
    return base_url if base_url.endswith("/responses") else f"{base_url}/responses"


def raw_responses_create(payload):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    ssl_context = ssl.create_default_context(cafile=certifi.where())
    request = Request(
        responses_endpoint_url(),
        data=body,
        headers={
            "Authorization": f"Bearer {ai_api_key()}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urlopen(request, timeout=120, context=ssl_context) as response:
            return json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Responses API 请求失败({exc.code}): {detail[:800]}") from exc


def responses_create(payload):
    client = ai_client()
    if client and hasattr(client, "responses"):
        return client.responses.create(**payload)
    return raw_responses_create(payload)


def chat_json(system_prompt, user_payload):
    payload_text = json.dumps(user_payload, ensure_ascii=False)
    if ai_api_type() == "responses":
        kwargs = {
            "model": ai_model(),
            "input": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": payload_text},
            ],
            "temperature": 0.2,
        }
        try:
            response = responses_create({**kwargs, "text": {"format": {"type": "json_object"}}})
        except Exception as exc:
            message = str(exc).lower()
            if "text" not in message and "format" not in message and "json_object" not in message:
                raise
            response = responses_create(kwargs)
        content = response_output_text(response)
        return parse_json_content(content)

    client = ai_client()
    if not client:
        return None
    response = client.chat.completions.create(
        model=ai_model(),
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": payload_text},
        ],
        response_format={"type": "json_object"},
        temperature=0.2,
    )
    return parse_json_content(response.choices[0].message.content)


def build_question_bank(raw_text, filename, use_ai=True):
    if not use_ai:
        return [normalize_question_item(item) for item in fallback_build_question_bank(raw_text)]
    system_prompt = (
        "你是题库生成器，只返回 JSON，不要 Markdown，不要解释。"
        "输出格式：{\"questions\":[{\"type\":\"choice|judge|fill|short\",\"prompt\":\"题干\",\"answer\":\"标准答案\",\"choices\":[\"A\",\"B\"],\"explanation\":\"简短解析\",\"source_text\":\"原文依据\"}]}. "
        "必须尽量把内容分成三类：choice 题必须有 2-4 个选项且只有一个标准答案；judge 题的 choices 必须是 [\"对\",\"错\"] 且 answer 只能是 对 或 错；"
        "fill/short 题的 choices 为空数组。题目要清晰，答案必须来自原文。"
    )
    data = chat_json(system_prompt, {"filename": filename, "text": " ".join(raw_text.split())[:12000]})
    if not data:
        return [normalize_question_item(item) for item in fallback_build_question_bank(raw_text)]
    raw_questions = data.get("questions", []) or fallback_build_question_bank(raw_text)
    normalized = [normalize_question_item(item) for item in raw_questions if isinstance(item, dict)]
    return normalized or [normalize_question_item(item) for item in fallback_build_question_bank(raw_text)]


def score_status(score):
    if score >= 0.85:
        return "correct"
    if score >= 0.4:
        return "partial"
    return "wrong"


def fallback_judge_answer(correct_answer, user_answer):
    is_match = normalize(correct_answer) in normalize(user_answer) or normalize(user_answer) in normalize(correct_answer)
    score = 1 if is_match else 0
    return {
        "is_correct": is_match,
        "status": score_status(score),
        "score": score,
        "feedback": "答案匹配。" if is_match else "与标准答案不完全一致。",
    }


def judge_answer(question_row, user_answer):
    qtype = question_row["qtype"]
    correct = question_row["answer"]
    if qtype == "judge":
        user_norm = normalize(user_answer)
        correct_norm = normalize(correct)
        if user_norm.startswith("对") or user_norm in {"yes", "true", "1"}:
            user_answer = "对"
        elif user_norm.startswith("错") or user_norm in {"no", "false", "0"}:
            user_answer = "错"
        return fallback_judge_answer(correct, user_answer)
    if qtype == "choice":
        return fallback_judge_answer(correct, user_answer)
    if qtype == "fill":
        return fallback_judge_answer(correct, user_answer)

    client = ai_client()
    if not client:
        return fallback_judge_answer(correct, user_answer)
    system_prompt = (
        "你是严谨但友好的判题器，只返回 JSON。"
        "输出格式：{\"score\":0.0-1.0,\"feedback\":\"简短反馈\"}。"
        "允许同义表达、顺序差异和部分正确。核心意思对给高分，部分对给中等分。"
    )
    try:
        data = chat_json(
            system_prompt,
            {"question": {"prompt": question_row["prompt"], "answer": correct}, "user_answer": user_answer},
        )
        if not data:
            return fallback_judge_answer(correct, user_answer)
        score = max(0, min(1, float(data.get("score", 0))))
        return {
            "is_correct": score >= 0.85,
            "status": score_status(score),
            "score": score,
            "feedback": data.get("feedback", ""),
        }
    except Exception:
        return fallback_judge_answer(correct, user_answer)


def question_to_json(row):
    return {
        "id": row["id"],
        "paperId": row["paper_id"],
        "type": row["qtype"],
        "prompt": row["prompt"],
        "answer": row["answer"],
        "choices": json.loads(row["choices_json"] or "[]"),
        "explanation": row["explanation"],
        "fullText": row["full_text"],
    }


@app.before_request
def _ensure_db():
    init_db()
    bootstrap_fixed_admin()


@app.get("/")
def home():
    return render_template("index.html")


@app.get("/api/me")
def me():
    identity = current_identity()
    return jsonify(
        {
            "isAdmin": bool(session.get("admin_id")),
            "isOwner": bool(session.get("is_owner")),
            "isUser": bool(session.get("user_id")),
            "username": identity["username"] if identity else None,
            "ai": {
                "enabled": ai_enabled(),
                "model": ai_model(),
                "apiType": ai_api_type(),
                "baseUrl": get_setting("ai_api_base", DEFAULT_AI_API_BASE) or "default",
            },
        }
    )


@app.post("/api/login")
def unified_login():
    data = request.get_json(force=True)
    username = data.get("username", "").strip()
    password = data.get("password", "")
    with get_db() as db:
        user = db.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["user_username"] = user["username"]
            return jsonify({"ok": True, "kind": "user", "username": user["username"]})
        admin = db.execute("SELECT * FROM admins WHERE username = ?", (username,)).fetchone()
        if admin and check_password_hash(admin["password_hash"], password):
            session.clear()
            session["admin_id"] = admin["id"]
            session["admin_username"] = admin["username"]
            session["is_owner"] = bool(admin["is_owner"])
            return jsonify({"ok": True, "kind": "admin", "username": admin["username"], "isOwner": bool(admin["is_owner"])})
    return jsonify({"error": "invalid_credentials"}), 401


@app.post("/api/auth/register")
def user_register():
    data = request.get_json(force=True)
    username = data.get("username", "").strip()
    password = data.get("password", "")
    if len(username) < 3 or len(password) < 8:
        return jsonify({"error": "用户名至少 3 位，密码至少 8 位"}), 400
    with get_db() as db:
        try:
            db.execute(
                "INSERT INTO users (username, password_hash, created_at) VALUES (?, ?, ?)",
                (username, generate_password_hash(password), now_iso()),
            )
            db.commit()
        except sqlite3.IntegrityError:
            return jsonify({"error": "用户名已存在"}), 409
    return jsonify({"ok": True})


@app.post("/api/auth/login")
def user_login():
    return unified_login()


@app.post("/api/auth/logout")
def user_logout():
    session.clear()
    return jsonify({"ok": True})


@app.post("/api/admin/login")
def admin_login():
    return unified_login()


@app.post("/api/admin/logout")
def admin_logout():
    session.clear()
    return jsonify({"ok": True})


@app.post("/api/logout")
def logout():
    session.clear()
    return jsonify({"ok": True})


@app.post("/api/admin/password")
@admin_required
def admin_password():
    data = request.get_json(force=True)
    old_password = data.get("oldPassword", "")
    new_password = data.get("newPassword", "")
    if len(new_password) < 8:
        return jsonify({"error": "新密码至少 8 位"}), 400
    with get_db() as db:
        admin = db.execute("SELECT * FROM admins WHERE id = ?", (session.get("admin_id"),)).fetchone()
        if not admin or not check_password_hash(admin["password_hash"], old_password):
            return jsonify({"error": "旧密码不正确"}), 401
        db.execute("UPDATE admins SET password_hash = ? WHERE id = ?", (generate_password_hash(new_password), session.get("admin_id")))
        db.commit()
    return jsonify({"ok": True})


@app.get("/api/admin/settings")
@owner_required
def admin_settings():
    return jsonify(
        {
            "fixedAdminUser": FIXED_ADMIN_USERNAME,
            "ownerAdminId": get_setting("owner_admin_id", ""),
            "aiApiKey": get_setting("ai_api_key", ""),
            "aiApiBase": get_setting("ai_api_base", ""),
            "aiModel": get_setting("ai_model", DEFAULT_AI_MODEL),
            "aiApiType": ai_api_type(),
        }
    )


@app.post("/api/admin/settings")
@owner_required
def save_admin_settings():
    data = request.get_json(force=True)
    fixed_admin_user = str(data.get("fixedAdminUser", "")).strip() or FIXED_ADMIN_USERNAME
    ai_api_key = str(data.get("aiApiKey", "")).strip()
    ai_api_base = str(data.get("aiApiBase", "")).strip()
    ai_model_value = str(data.get("aiModel", "")).strip() or DEFAULT_AI_MODEL
    ai_api_type_value = str(data.get("aiApiType", "chat")).strip().lower()
    if ai_api_type_value not in {"chat", "responses"}:
        ai_api_type_value = "chat"
    with get_db() as db:
        admin = db.execute("SELECT id FROM admins WHERE username = ?", (fixed_admin_user,)).fetchone()
        if not admin:
            return jsonify({"error": "fixed_admin_not_found"}), 404
        db.execute("UPDATE admins SET is_owner = CASE WHEN id = ? THEN 1 ELSE 0 END", (admin["id"],))
        db.execute("INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value", ("owner_admin_id", str(admin["id"])))
        db.execute("INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value", ("ai_api_key", ai_api_key))
        db.execute("INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value", ("ai_api_base", ai_api_base))
        db.execute("INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value", ("ai_model", ai_model_value))
        db.execute("INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value", ("ai_api_type", ai_api_type_value))
        db.commit()
    return jsonify({"ok": True})


@app.get("/api/admin/admins")
@owner_required
def list_admins():
    with get_db() as db:
        rows = db.execute("SELECT id, username, is_owner, created_at FROM admins ORDER BY is_owner DESC, id ASC").fetchall()
    return jsonify([dict(row) for row in rows])


@app.post("/api/admin/admins")
@owner_required
def create_admin():
    data = request.get_json(force=True)
    username = data.get("username", "").strip()
    password = data.get("password", "")
    if len(username) < 3 or len(password) < 8:
        return jsonify({"error": "用户名至少 3 位，密码至少 8 位"}), 400
    with get_db() as db:
        try:
            db.execute(
                "INSERT INTO admins (username, password_hash, is_owner, created_at) VALUES (?, ?, 0, ?)",
                (username, generate_password_hash(password), now_iso()),
            )
            db.commit()
        except sqlite3.IntegrityError:
            return jsonify({"error": "管理员账号已存在"}), 409
    return jsonify({"ok": True})


@app.delete("/api/admin/admins/<int:admin_id>")
@owner_required
def delete_admin(admin_id):
    with get_db() as db:
        row = db.execute("SELECT is_owner FROM admins WHERE id = ?", (admin_id,)).fetchone()
        if not row:
            return jsonify({"error": "not_found"}), 404
        if row["is_owner"]:
            return jsonify({"error": "固定管理员不能删除"}), 400
        db.execute("DELETE FROM admins WHERE id = ?", (admin_id,))
        db.commit()
    return jsonify({"ok": True})


@app.post("/api/admin/upload-paper")
@admin_required
def upload_paper():
    if "file" not in request.files:
        return jsonify({"error": "missing_file"}), 400
    file = request.files["file"]
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in {".pdf", ".txt", ".text", ".md", ".markdown", ".docx"}:
        return jsonify({"error": "unsupported_file_type"}), 400

    stored_filename, file_path = save_uploaded_file(file)
    raw_text = parse_uploaded_file(file_path, file.filename)
    use_ai = request.form.get("useAi", "true").lower() in {"1", "true", "yes", "on"}
    paper_title = request.form.get("paperTitle", "").strip() or os.path.splitext(file.filename)[0]
    questions = build_question_bank(raw_text, file.filename, use_ai=use_ai)
    created_at = now_iso()

    with get_db() as db:
        paper_id = db.execute(
            "INSERT INTO papers (title, file_type, filename, stored_filename, raw_text, use_ai, created_by, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (
                paper_title,
                file_type_label(file.filename),
                file.filename,
                stored_filename,
                raw_text,
                int(use_ai),
                session.get("admin_id"),
                created_at,
            ),
        ).lastrowid
        created = 0
        for item in questions:
            prompt = str(item.get("prompt", "")).strip()
            answer = str(item.get("answer", "")).strip()
            if not prompt or not answer:
                continue
            db.execute(
                """
                INSERT INTO questions (paper_id, qtype, prompt, answer, choices_json, explanation, full_text, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    paper_id,
                    item.get("type", "fill"),
                    prompt,
                    answer,
                    json.dumps(item.get("choices", []), ensure_ascii=False),
                    item.get("explanation", ""),
                    item.get("source_text", "") or raw_text[:500],
                    created_at,
                ),
            )
            created += 1
        db.commit()

    return jsonify({"ok": True, "paperId": paper_id, "questionsCreated": created, "aiEnabled": bool(use_ai and ai_enabled())})


@app.post("/api/admin/upload-pdf")
@admin_required
def upload_paper_legacy():
    return upload_paper()


@app.post("/api/admin/papers")
@admin_required
def create_manual_paper():
    data = request.get_json(force=True)
    title = str(data.get("title", "")).strip()
    if not title:
        return jsonify({"error": "missing_title"}), 400
    created_at = now_iso()
    with get_db() as db:
        paper_id = db.execute(
            "INSERT INTO papers (title, file_type, filename, stored_filename, raw_text, use_ai, created_by, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (title, "手动", title, "", "", 0, session.get("admin_id"), created_at),
        ).lastrowid
        db.commit()
    return jsonify({"ok": True, "paperId": paper_id})


@app.post("/api/admin/papers/<int:paper_id>/questions")
@admin_required
def create_manual_question(paper_id):
    data = request.get_json(force=True)
    item = normalize_question_item(
        {
            "type": data.get("type", "fill"),
            "prompt": data.get("prompt", ""),
            "answer": data.get("answer", ""),
            "choices": data.get("choices", []),
            "explanation": data.get("explanation", ""),
            "source_text": data.get("sourceText", ""),
        }
    )
    if not item["prompt"] or not item["answer"]:
        return jsonify({"error": "missing_question_or_answer"}), 400
    with get_db() as db:
        paper = db.execute("SELECT id FROM papers WHERE id = ?", (paper_id,)).fetchone()
        if not paper:
            return jsonify({"error": "paper_not_found"}), 404
        question_id = db.execute(
            """
            INSERT INTO questions (paper_id, qtype, prompt, answer, choices_json, explanation, full_text, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                paper_id,
                item["type"],
                item["prompt"],
                item["answer"],
                json.dumps(item["choices"], ensure_ascii=False),
                item["explanation"],
                item["source_text"] or item["prompt"],
                now_iso(),
            ),
        ).lastrowid
        db.commit()
    return jsonify({"ok": True, "questionId": question_id})


@app.delete("/api/admin/papers/<int:paper_id>")
@admin_required
def delete_paper(paper_id):
    with get_db() as db:
        paper = db.execute("SELECT stored_filename FROM papers WHERE id = ?", (paper_id,)).fetchone()
        if not paper:
            return jsonify({"error": "not_found"}), 404
        question_ids = [row["id"] for row in db.execute("SELECT id FROM questions WHERE paper_id = ?", (paper_id,)).fetchall()]
        if question_ids:
            placeholders = ",".join("?" for _ in question_ids)
            db.execute(f"DELETE FROM attempts WHERE question_id IN ({placeholders})", question_ids)
            db.execute(f"DELETE FROM questions WHERE id IN ({placeholders})", question_ids)
        db.execute("DELETE FROM papers WHERE id = ?", (paper_id,))
        db.commit()
    if paper["stored_filename"]:
        try:
            path = os.path.join(UPLOAD_DIR, paper["stored_filename"])
            if os.path.exists(path):
                os.remove(path)
        except OSError:
            pass
    return jsonify({"ok": True})


@app.post("/api/admin/test-ai")
@owner_required
def test_ai():
    if not ai_enabled():
        return jsonify({"ok": False, "enabled": False, "message": "ai_not_configured"}), 400
    try:
        result = chat_json(
            "只返回 JSON，不要 Markdown。",
            {"task": "请返回 {\"ok\":true,\"message\":\"AI通了\"}"},
        )
        return jsonify({"ok": True, "enabled": True, "result": result or {}, "model": ai_model(), "apiType": ai_api_type()})
    except Exception as exc:
        return jsonify({"ok": False, "enabled": True, "error": str(exc)}), 500


@app.get("/api/papers")
def list_papers():
    with get_db() as db:
        rows = db.execute(
            """
            SELECT p.id, p.title, p.file_type, p.filename, p.stored_filename, p.use_ai, p.created_at,
                   COUNT(q.id) AS question_count
            FROM papers p
            LEFT JOIN questions q ON q.paper_id = p.id
            GROUP BY p.id
            ORDER BY p.id DESC
            """
        ).fetchall()
    return jsonify(
        [
            {
                "id": row["id"],
                "title": row["title"],
                "fileType": row["file_type"],
                "filename": row["filename"],
                "createdAt": row["created_at"],
                "questionCount": row["question_count"],
                "useAi": bool(row["use_ai"]),
                "url": f"/uploads/{row['stored_filename']}" if row["stored_filename"] else None,
            }
            for row in rows
        ]
    )


@app.get("/api/sources")
def list_sources():
    return list_papers()


@app.get("/uploads/<path:stored_filename>")
def uploaded_file(stored_filename):
    if "/" in stored_filename or "\\" in stored_filename:
        abort(404)
    return send_from_directory(UPLOAD_DIR, stored_filename, as_attachment=False)


@app.get("/api/questions")
def list_questions():
    paper_id = request.args.get("paperId")
    limit = min(int(request.args.get("limit", 100)), 300)
    params = []
    sql = "SELECT id, paper_id, qtype, prompt, answer, choices_json, explanation, full_text FROM questions"
    if paper_id:
        sql += " WHERE paper_id = ?"
        params.append(int(paper_id))
    sql += " ORDER BY id DESC LIMIT ?"
    params.append(limit)
    with get_db() as db:
        rows = db.execute(sql, tuple(params)).fetchall()
    return jsonify([question_to_json(row) for row in rows])


@app.get("/api/questions/random")
def random_question():
    paper_id = request.args.get("paperId")
    params = []
    sql = "SELECT id, paper_id, qtype, prompt, answer, choices_json, explanation, full_text FROM questions"
    if paper_id:
        sql += " WHERE paper_id = ?"
        params.append(int(paper_id))
    if session.get("user_id"):
        sql += """
            AND id NOT IN (
                SELECT question_id FROM attempts
                WHERE user_id = ? AND score >= 1
            )
        """ if paper_id else """
            WHERE id NOT IN (
                SELECT question_id FROM attempts
                WHERE user_id = ? AND score >= 1
            )
        """
        params.append(session.get("user_id"))
    sql += " ORDER BY RANDOM() LIMIT 1"
    with get_db() as db:
        row = db.execute(sql, tuple(params)).fetchone()
    if not row:
        return jsonify({"error": "no_questions"}), 404
    return jsonify(question_to_json(row))


@app.get("/api/progress")
@login_required
def progress():
    paper_id = request.args.get("paperId")
    params = [session.get("user_id") or 0]
    sql = """
        SELECT COUNT(DISTINCT a.question_id) AS correct_count
        FROM attempts a
        JOIN questions q ON q.id = a.question_id
        WHERE a.user_id = ? AND a.score >= 1
    """
    if paper_id:
        sql += " AND q.paper_id = ?"
        params.append(int(paper_id))
    with get_db() as db:
        row = db.execute(sql, tuple(params)).fetchone()
    return jsonify({"correctCount": row["correct_count"] if row else 0})


@app.get("/api/papers/<int:paper_id>/participation")
@login_required
def get_participation(paper_id):
    with get_db() as db:
        row = db.execute(
            "SELECT participate FROM paper_participants WHERE paper_id = ? AND user_id = ?",
            (paper_id, session.get("user_id") or 0),
        ).fetchone()
    return jsonify({"selected": row is not None, "participate": bool(row["participate"]) if row else False})


@app.post("/api/papers/<int:paper_id>/participation")
@login_required
def set_participation(paper_id):
    data = request.get_json(force=True)
    participate = 1 if bool(data.get("participate")) else 0
    now = now_iso()
    with get_db() as db:
        paper = db.execute("SELECT id FROM papers WHERE id = ?", (paper_id,)).fetchone()
        if not paper:
            return jsonify({"error": "paper_not_found"}), 404
        db.execute(
            """
            INSERT INTO paper_participants (paper_id, user_id, participate, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT(paper_id, user_id) DO UPDATE SET participate = excluded.participate, updated_at = excluded.updated_at
            """,
            (paper_id, session.get("user_id") or 0, participate, now, now),
        )
        db.commit()
    return jsonify({"ok": True, "participate": bool(participate)})


@app.get("/api/papers/<int:paper_id>/leaderboard")
def leaderboard(paper_id):
    with get_db() as db:
        rows = db.execute(
            """
            SELECT u.username,
                   COUNT(DISTINCT CASE WHEN a.score >= 1 THEN a.question_id END) AS correct_count,
                   COUNT(DISTINCT q.id) AS total_count,
                   MAX(a.created_at) AS last_answer_at
            FROM paper_participants pp
            JOIN users u ON u.id = pp.user_id
            JOIN questions q ON q.paper_id = pp.paper_id
            LEFT JOIN attempts a ON a.question_id = q.id AND a.user_id = pp.user_id AND a.participate = 1
            WHERE pp.paper_id = ? AND pp.participate = 1
            GROUP BY pp.user_id
            ORDER BY correct_count DESC, last_answer_at ASC
            LIMIT 50
            """,
            (paper_id,),
        ).fetchall()
    return jsonify(
        [
            {
                "rank": index + 1,
                "username": row["username"],
                "correctCount": row["correct_count"],
                "totalCount": row["total_count"],
                "lastAnswerAt": row["last_answer_at"],
            }
            for index, row in enumerate(rows)
        ]
    )


@app.post("/api/attempts")
@login_required
def record_attempt():
    data = request.get_json(force=True)
    question_id = int(data["questionId"])
    user_answer = data.get("userAnswer", "")
    with get_db() as db:
        row = db.execute("SELECT qtype, prompt, answer, full_text FROM questions WHERE id = ?", (question_id,)).fetchone()
        if not row:
            return jsonify({"error": "question_not_found"}), 404
        judged = judge_answer(row, user_answer)
        score = max(0, min(1, float(judged.get("score", 0))))
        status = judged.get("status") or score_status(score)
        is_correct = int(status == "correct")
        participation = db.execute(
            """
            SELECT pp.participate
            FROM paper_participants pp
            JOIN questions q ON q.paper_id = pp.paper_id
            WHERE q.id = ? AND pp.user_id = ?
            """,
            (question_id, session.get("user_id") or 0),
        ).fetchone()
        participate = int(bool(participation and participation["participate"]))
        db.execute(
            "INSERT INTO attempts (question_id, user_id, user_answer, is_correct, score, feedback, created_at, participate) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (question_id, session.get("user_id") or 0, user_answer, is_correct, score, judged.get("feedback", ""), now_iso(), participate),
        )
        db.commit()
    return jsonify({"ok": True, "isCorrect": bool(is_correct), "status": status, "score": score, "answer": row["answer"], "feedback": judged.get("feedback", "")})


if __name__ == "__main__":
    init_db()
    bootstrap_fixed_admin()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 8000)), debug=True)
